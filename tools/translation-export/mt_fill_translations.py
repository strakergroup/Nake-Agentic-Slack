"""Fill missing translation workbook rows with Google Cloud Translation MT."""

from __future__ import annotations

import argparse
import glob
import html
import json
import os
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from export_missing_strings import ensure_repo_root_on_path
from sqlalchemy import text

TRANSLATION_COLUMN = "target_text"
DB_LABEL_COLUMN = "source_text"
DB_LANG_COLUMN = "target_language"
NOTES_COLUMN = "notes"
ENGLISH_PREFIXES = ("en", "gb", "us")
DEFAULT_GOOGLE_TRANSLATE_BATCH_SIZE = 100
GOOGLE_TRANSLATE_BATCH_SIZE_ENV = "TRANSLATION_EXPORT_MT_BATCH_SIZE"
GOOGLE_CLOUD_PROJECT_ENV = "GOOGLE_CLOUD_PROJECT"
GOOGLE_CLOUD_LOCATION_ENV = "GOOGLE_CLOUD_LOCATION"
GOOGLE_CLOUD_DEFAULT_LOCATION = "global"
GOOGLE_SOURCE_LANGUAGE_CODE = "en"
GOOGLE_TRANSLATE_MIME_TYPE = "text/html"
COLUMN_ALIASES = {
    TRANSLATION_COLUMN: ("translation",),
    DB_LABEL_COLUMN: ("db_label",),
    DB_LANG_COLUMN: ("db_lang",),
}


@dataclass(frozen=True)
class WorkbookTranslationRow:
    row_number: int
    source_text: str
    db_lang: str
    google_target_lang: str


def configured_google_project() -> str:
    project_id = os.getenv(GOOGLE_CLOUD_PROJECT_ENV, "").strip()
    if not project_id:
        raise ValueError(
            f"{GOOGLE_CLOUD_PROJECT_ENV} must be set for Google Cloud Translation"
        )
    return project_id


def configured_google_location() -> str:
    return (
        os.getenv(GOOGLE_CLOUD_LOCATION_ENV, GOOGLE_CLOUD_DEFAULT_LOCATION).strip()
        or GOOGLE_CLOUD_DEFAULT_LOCATION
    )


def configured_batch_size(batch_size: int | None = None) -> int:
    raw_value = (
        str(batch_size)
        if batch_size is not None
        else os.getenv(GOOGLE_TRANSLATE_BATCH_SIZE_ENV, "")
    ).strip()
    if not raw_value:
        return DEFAULT_GOOGLE_TRANSLATE_BATCH_SIZE
    try:
        parsed = int(raw_value)
    except ValueError as exc:
        raise ValueError(
            f"{GOOGLE_TRANSLATE_BATCH_SIZE_ENV} must be a positive integer"
        ) from exc
    if parsed < 1:
        raise ValueError(
            f"{GOOGLE_TRANSLATE_BATCH_SIZE_ENV} must be a positive integer"
        )
    return parsed


def build_google_translate_client() -> Any:
    try:
        from google.cloud import translate_v3
    except ImportError as exc:
        raise RuntimeError(
            "google-cloud-translate is required for MT fill; install project "
            "dependencies before running this tool"
        ) from exc

    return translate_v3.TranslationServiceClient()


def fetch_mt_language_map() -> dict[str, str]:
    ensure_repo_root_on_path()
    from app.database import engines

    with engines["translators_readonly"].connect() as conn:
        result = conn.execute(
            text(
                """
                SELECT bcp_47, shortname, site_shortname, lang, google_code
                FROM obj_m_langs
                """
            )
        )

        language_map: dict[str, str] = {}
        for row in result:
            mt_code = next(
                (
                    str(value).strip()
                    for value in (row.google_code, row.bcp_47, row.site_shortname)
                    if value and str(value).strip()
                ),
                "",
            )
            if not mt_code:
                continue

            for value in (
                row.shortname,
                row.site_shortname,
                row.bcp_47,
                row.lang,
                row.google_code,
            ):
                if value and str(value).strip():
                    language_map[str(value).strip().lower()] = mt_code
        return language_map


def resolve_mt_target_language(
    db_lang: str,
    mt_language_map: dict[str, str],
) -> str:
    return mt_language_map.get(db_lang.strip().lower(), db_lang)


def is_english_language(language_code: str) -> bool:
    return language_code.strip().lower().startswith(ENGLISH_PREFIXES)


def validate_mt_translation(
    source_text: str, target_lang: str, translation: str
) -> None:
    if is_english_language(target_lang):
        return
    if translation.strip() == source_text.strip():
        raise ValueError(
            f"Google returned unchanged source text for non-English target "
            f"{target_lang!r}"
        )


def validate_google_target_language(db_lang: str, google_target_lang: str) -> None:
    if not is_english_language(db_lang) and is_english_language(google_target_lang):
        raise ValueError(
            f"Google target language resolved to English fallback for non-English "
            f"DB language {db_lang!r}: {google_target_lang!r}"
        )


def has_suspicious_matching_source_batch(total_rows: int, matching_rows: int) -> bool:
    if total_rows < 5:
        return False
    return matching_rows / total_rows >= 0.8


def batched(items: list[WorkbookTranslationRow], batch_size: int):
    for start in range(0, len(items), batch_size):
        yield items[start : start + batch_size]


def google_translate_texts(
    texts: list[str],
    target_lang: str,
    client: Any | None = None,
    project_id: str | None = None,
    location: str | None = None,
) -> list[str]:
    if not texts:
        return []

    translate_client = client or build_google_translate_client()
    parent = (
        f"projects/{project_id or configured_google_project()}/"
        f"locations/{location or configured_google_location()}"
    )
    response = translate_client.translate_text(
        contents=texts,
        parent=parent,
        mime_type=GOOGLE_TRANSLATE_MIME_TYPE,
        source_language_code=GOOGLE_SOURCE_LANGUAGE_CODE,
        target_language_code=target_lang,
    )
    translations = [
        str(translation.translated_text) for translation in response.translations
    ]
    if len(translations) != len(texts):
        raise ValueError(
            f"Google returned {len(translations)} translation(s) for "
            f"{len(texts)} requested text(s) targeting {target_lang!r}"
        )
    return translations


def mt_translate_db_label(text: str, target_lang: str) -> str:
    return google_translate_texts([text], target_lang)[0]


def header_indexes(sheet) -> dict[str, int]:
    headers = [cell.value for cell in sheet[1]]
    return {
        str(header): index + 1
        for index, header in enumerate(headers)
        if isinstance(header, str)
    }


def column_index(indexes: dict[str, int], column: str) -> int | None:
    if column in indexes:
        return indexes[column]
    return next(
        (
            indexes[alias]
            for alias in COLUMN_ALIASES.get(column, ())
            if alias in indexes
        ),
        None,
    )


def append_note(existing: object, note: str) -> str:
    existing_text = "" if existing is None else str(existing).strip()
    if not existing_text:
        return note
    return f"{existing_text}; {note}"


def decode_html_entities(value: str) -> str:
    return html.unescape(value)


def fill_workbook_translations(
    xlsx_path: Path,
    mt_language_map: dict[str, str] | None = None,
    batch_size: int | None = None,
    google_client: Any | None = None,
) -> tuple[int, int]:
    if mt_language_map is None:
        mt_language_map = fetch_mt_language_map()
    resolved_batch_size = configured_batch_size(batch_size)
    translate_client = google_client
    workbook = openpyxl.load_workbook(xlsx_path)
    try:
        sheet = workbook.active
        indexes = header_indexes(sheet)
        label_index = column_index(indexes, DB_LABEL_COLUMN)
        lang_index = column_index(indexes, DB_LANG_COLUMN)
        translation_index = column_index(indexes, TRANSLATION_COLUMN)
        missing_columns = [
            column
            for column, index in (
                (DB_LABEL_COLUMN, label_index),
                (DB_LANG_COLUMN, lang_index),
                (TRANSLATION_COLUMN, translation_index),
            )
            if index is None
        ]
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise ValueError(f"Workbook is missing required columns: {missing}")
        assert label_index is not None
        assert lang_index is not None
        assert translation_index is not None

        notes_index = indexes.get(NOTES_COLUMN)
        if notes_index is None:
            notes_index = sheet.max_column + 1
            sheet.cell(row=1, column=notes_index).value = NOTES_COLUMN
        filled = 0
        total = 0
        errors: list[str] = []
        row_counts_by_lang: dict[str, int] = defaultdict(int)
        matching_source_rows_by_lang: dict[str, list[int]] = defaultdict(list)
        rows_by_google_lang: dict[str, list[WorkbookTranslationRow]] = defaultdict(list)
        for row_number in range(2, sheet.max_row + 1):
            label_cell = sheet.cell(row=row_number, column=label_index)
            lang_cell = sheet.cell(row=row_number, column=lang_index)
            translation_cell = sheet.cell(
                row=row_number,
                column=translation_index,
            )
            if not label_cell.value or not lang_cell.value:
                continue
            db_lang = str(lang_cell.value)
            source_text = str(label_cell.value)
            total += 1
            if not is_english_language(db_lang):
                row_counts_by_lang[db_lang] += 1
            if translation_cell.value and str(translation_cell.value).strip():
                if (
                    not is_english_language(db_lang)
                    and str(translation_cell.value).strip() == source_text.strip()
                ):
                    matching_source_rows_by_lang[db_lang].append(row_number)
                continue
            try:
                google_target_lang = resolve_mt_target_language(
                    db_lang,
                    mt_language_map,
                )
                validate_google_target_language(db_lang, google_target_lang)
                rows_by_google_lang[google_target_lang].append(
                    WorkbookTranslationRow(
                        row_number=row_number,
                        source_text=source_text,
                        db_lang=db_lang,
                        google_target_lang=google_target_lang,
                    )
                )
            except Exception as exc:
                errors.append(f"row {row_number}: {exc}")
                notes_cell = sheet.cell(row=row_number, column=notes_index)
                notes_cell.value = append_note(notes_cell.value, f"MT error: {exc}")

        for google_target_lang, rows in rows_by_google_lang.items():
            for batch in batched(rows, resolved_batch_size):
                try:
                    if translate_client is None:
                        translate_client = build_google_translate_client()
                    translations = google_translate_texts(
                        [row.source_text for row in batch],
                        google_target_lang,
                        client=translate_client,
                    )
                except Exception as exc:
                    for row in batch:
                        errors.append(f"row {row.row_number}: {exc}")
                        notes_cell = sheet.cell(
                            row=row.row_number,
                            column=notes_index,
                        )
                        notes_cell.value = append_note(
                            notes_cell.value,
                            f"MT error: {exc}",
                        )
                    continue

                for row, translation in zip(batch, translations, strict=True):
                    translation_text = decode_html_entities(translation)
                    translation_cell = sheet.cell(
                        row=row.row_number,
                        column=translation_index,
                    )
                    translation_cell.value = translation_text
                    if (
                        not is_english_language(row.db_lang)
                        and translation_text.strip() == row.source_text.strip()
                    ):
                        matching_source_rows_by_lang[row.db_lang].append(row.row_number)
                    filled += 1

        for db_lang, matching_rows in matching_source_rows_by_lang.items():
            if has_suspicious_matching_source_batch(
                row_counts_by_lang[db_lang], len(matching_rows)
            ):
                error = (
                    f"MT output matched source text for {len(matching_rows)} of "
                    f"{row_counts_by_lang[db_lang]} non-English {db_lang!r} row(s)"
                )
                errors.append(error)
                for row_number in matching_rows:
                    notes_cell = sheet.cell(row=row_number, column=notes_index)
                    notes_cell.value = append_note(
                        notes_cell.value,
                        "MT validation error: translation matches source text across "
                        "a suspicious share of this workbook",
                    )

        workbook.save(xlsx_path)
        if errors:
            preview = "; ".join(errors[:3])
            more = "" if len(errors) <= 3 else f"; +{len(errors) - 3} more"
            raise ValueError(
                f"MT fill failed for {len(errors)} row(s) in {xlsx_path}: "
                f"{preview}{more}"
            )
        return filled, total
    finally:
        workbook.close()


def resolve_workbook_paths(input_pattern: str) -> list[Path]:
    matches = sorted(glob.glob(input_pattern))
    if matches:
        return [Path(match) for match in matches]
    return [Path(input_pattern)]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Fill translation-export workbook rows with Google Cloud MT."
    )
    parser.add_argument(
        "--input",
        default=str(Path(__file__).parent / "output" / "missing_strings.xlsx"),
        help="Workbook generated by export_missing_strings.py. Globs are supported.",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=None,
        help=(
            "Maximum workbook rows to translate per Google request. Defaults to "
            f"{GOOGLE_TRANSLATE_BATCH_SIZE_ENV} or "
            f"{DEFAULT_GOOGLE_TRANSLATE_BATCH_SIZE}."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    summary = []
    for input_path in resolve_workbook_paths(args.input):
        filled, total = fill_workbook_translations(
            input_path,
            batch_size=args.batch_size,
        )
        summary.append({"file": str(input_path), "filled": filled, "total": total})
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    sys.exit(main())
