"""Export app UI strings that are missing from the translation database."""

from __future__ import annotations

import argparse
import ast
import csv
import json
import os
import re
import sys
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

from sqlalchemy import text

DEFAULT_SLACK_LOCALES = (
    "de-DE",
    "es-ES",
    "es-LA",
    "fr-FR",
    "fr-CA",
    "it-IT",
    "ja-JP",
    "ko-KR",
    "pt-BR",
    "zh-CN",
    "zh-TW",
)
ENGLISH_PREFIXES = ("en", "gb", "us")
PLACEHOLDER_PATTERN = re.compile(r":\w+:|\{.*?\}")
X_TAG_PATTERN = re.compile(r'<x id="?(\d+)"?\s*/?>')
OUTPUT_COLUMNS = (
    "source_language",
    "target_language",
    "source_text",
    "target_text",
    "max_length",
)
TRANSLATOR_METADATA_SHEET = "_translation_metadata"
TRANSLATOR_METADATA_COLUMNS = ("source_text", "max_length")


@dataclass
class StringEntry:
    source_text: str
    db_label: str
    max_length: int = 0
    locations: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class LocaleTarget:
    slack_locale: str
    db_lang: str
    is_english: bool
    resolved_from_db: bool


@dataclass(frozen=True)
class MissingStringRow:
    source_language: str
    target_language: str
    source_text: str
    target_text: str
    max_length: int


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def ensure_repo_root_on_path() -> None:
    root = str(repo_root())
    if root not in sys.path:
        sys.path.insert(0, root)


def parse_languages(value: str | None) -> list[str]:
    """Parse comma-separated Slack locale codes, preserving order."""
    if not value:
        return list(DEFAULT_SLACK_LOCALES)

    languages: list[str] = []
    seen: set[str] = set()
    for raw_language in value.split(","):
        language = raw_language.strip()
        key = language.lower()
        if not language or key in seen:
            continue
        languages.append(language)
        seen.add(key)
    return languages


def tag_placeholders(text_value: str) -> str:
    """Match app.translate.Translator placeholder tagging for DB labels."""
    tagged_text = text_value
    for index, match in enumerate(PLACEHOLDER_PATTERN.finditer(text_value), start=1):
        tagged_text = tagged_text.replace(match.group(), f'<x id="{index}"/>')
    return tagged_text


def normalize_db_label_for_lookup(label: str) -> str:
    """Mirror MySQL label equality and treat x-tag styles as equivalent."""
    normalized_label = X_TAG_PATTERN.sub(r'<x id="\1"/>', label)
    return normalized_label.rstrip(" ").casefold()


def extract_max_length(node: ast.Call) -> int:
    if len(node.args) > 1 and isinstance(node.args[1], ast.Constant):
        value = node.args[1].value
        if isinstance(value, int):
            return value

    for keyword in node.keywords:
        if keyword.arg == "max_length" and isinstance(keyword.value, ast.Constant):
            value = keyword.value.value
            if isinstance(value, int):
                return value

    return 0


def extract_entries_from_file(path: Path, root: Path) -> list[StringEntry]:
    source = path.read_text(encoding="utf-8")
    tree = ast.parse(source, filename=str(path))
    entries: list[StringEntry] = []

    for node in ast.walk(tree):
        if not (
            isinstance(node, ast.Call)
            and isinstance(node.func, ast.Name)
            and node.func.id == "_"
            and node.args
        ):
            continue

        first_arg = node.args[0]
        if not (
            isinstance(first_arg, ast.Constant) and isinstance(first_arg.value, str)
        ):
            continue

        source_text = first_arg.value
        entries.append(
            StringEntry(
                source_text=source_text,
                db_label=tag_placeholders(source_text),
                max_length=extract_max_length(node),
                locations=[f"{path.relative_to(root)}:{node.lineno}"],
            )
        )

    return entries


def collect_string_entries(
    source_dir: Path, root: Path | None = None
) -> list[StringEntry]:
    root = root or source_dir
    unique_entries: dict[tuple[str, int], StringEntry] = {}

    for path in sorted(source_dir.rglob("*.py")):
        for entry in extract_entries_from_file(path, root):
            key = (entry.source_text, entry.max_length)
            if key in unique_entries:
                unique_entries[key].locations.extend(entry.locations)
            else:
                unique_entries[key] = entry

    return list(unique_entries.values())


def resolve_locale_target(
    slack_locale: str,
    language_map: dict[str, str],
) -> LocaleTarget:
    db_lang = language_map.get(slack_locale.lower(), slack_locale)
    is_english = db_lang.lower().startswith(ENGLISH_PREFIXES)
    return LocaleTarget(
        slack_locale=slack_locale,
        db_lang=db_lang,
        is_english=is_english,
        resolved_from_db=slack_locale.lower() in language_map,
    )


def build_missing_rows(
    entries: Sequence[StringEntry],
    slack_locales: Sequence[str],
    language_map: dict[str, str],
    existing_labels_by_lang: dict[str, set[str]],
    include_english: bool = False,
) -> list[MissingStringRow]:
    rows: list[MissingStringRow] = []

    for slack_locale in slack_locales:
        target = resolve_locale_target(slack_locale, language_map)
        if target.is_english and not include_english:
            continue

        existing_labels = {
            normalize_db_label_for_lookup(label)
            for label in existing_labels_by_lang.get(target.db_lang, set())
        }
        for entry in entries:
            if normalize_db_label_for_lookup(entry.db_label) in existing_labels:
                continue
            rows.append(
                MissingStringRow(
                    source_language="en",
                    target_language=target.db_lang,
                    source_text=entry.db_label,
                    target_text="",
                    max_length=entry.max_length,
                )
            )

    return rows


def fetch_language_map() -> dict[str, str]:
    ensure_repo_root_on_path()
    from app.database import engines

    with engines["translators_readonly"].connect() as conn:
        result = conn.execute(
            text(
                """
                SELECT bcp_47, shortname
                FROM obj_m_langs
                WHERE bcp_47 IS NOT NULL
                AND bcp_47 != ''
                """
            )
        )
        return {str(row[0]).lower(): str(row[1]) for row in result}


def fetch_existing_labels(db_langs: Iterable[str]) -> dict[str, set[str]]:
    ensure_repo_root_on_path()
    from app.database import engines

    labels_by_lang: dict[str, set[str]] = {}
    with engines["sitemanager_readonly"].connect() as conn:
        for db_lang in sorted(set(db_langs)):
            result = conn.execute(
                text(
                    """
                    SELECT DISTINCT label
                    FROM obj_stringtranslator
                    WHERE lang = :lang
                    """
                ).bindparams(lang=db_lang)
            )
            labels_by_lang[db_lang] = {str(row[0]) for row in result}
    return labels_by_lang


def write_csv(rows: Sequence[MissingStringRow], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(asdict(row) for row in rows)


def write_translator_csv(
    rows: Sequence[MissingStringRow],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerows([row.source_text] for row in rows)


def write_json(rows: Sequence[MissingStringRow], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "missing_count": len(rows),
        "rows": [asdict(row) for row in rows],
    }
    output_path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )


def write_xlsx(rows: Sequence[MissingStringRow], output_path: Path) -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "XLSX export requires openpyxl. Run through the dev environment."
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Missing Strings"
    sheet.append(list(OUTPUT_COLUMNS))
    for row in rows:
        sheet.append([getattr(row, column) for column in OUTPUT_COLUMNS])
    workbook.save(output_path)


def write_translator_xlsx(
    rows: Sequence[MissingStringRow],
    output_path: Path,
) -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "XLSX export requires openpyxl. Run through the dev environment."
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Translations"
    for row_number, row in enumerate(rows, start=1):
        sheet.cell(row=row_number, column=1).value = row.source_text

    metadata_sheet = workbook.create_sheet(TRANSLATOR_METADATA_SHEET)
    metadata_sheet.append(list(TRANSLATOR_METADATA_COLUMNS))
    for row in rows:
        metadata_sheet.append([row.source_text, row.max_length])
    metadata_sheet.sheet_state = "hidden"
    workbook.save(output_path)


def write_rows(
    rows: Sequence[MissingStringRow],
    output_path: Path,
    output_format: str,
    audience: str = "internal",
) -> None:
    if audience == "translator":
        if output_format == "csv":
            write_translator_csv(rows, output_path)
            return
        if output_format == "xlsx":
            write_translator_xlsx(rows, output_path)
            return
        raise ValueError("Translator exports support only csv and xlsx formats")

    if output_format == "csv":
        write_csv(rows, output_path)
        return
    if output_format == "json":
        write_json(rows, output_path)
        return
    if output_format == "xlsx":
        write_xlsx(rows, output_path)
        return
    raise ValueError(f"Unsupported output format: {output_format}")


def slugify_language(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_") or "unknown"


def split_output_path(output_path: Path, db_lang: str, output_format: str) -> Path:
    suffix = output_path.suffix or f".{output_format}"
    stem = output_path.stem if output_path.suffix else output_path.name
    return output_path.with_name(f"{stem}_{slugify_language(db_lang)}{suffix}")


def write_rows_by_language(
    rows: Sequence[MissingStringRow],
    output_path: Path,
    output_format: str,
    db_langs: Sequence[str],
    audience: str = "internal",
) -> list[Path]:
    rows_by_lang: dict[str, list[MissingStringRow]] = defaultdict(list)
    for row in rows:
        rows_by_lang[row.target_language].append(row)

    output_paths: list[Path] = []
    for db_lang in dict.fromkeys(db_langs):
        language_output_path = split_output_path(output_path, db_lang, output_format)
        write_rows(
            rows_by_lang.get(db_lang, []),
            language_output_path,
            output_format,
            audience,
        )
        output_paths.append(language_output_path)
    return output_paths


def infer_format(output_path: Path, explicit_format: str | None) -> str:
    if explicit_format:
        return explicit_format
    suffix = output_path.suffix.lower().lstrip(".")
    return suffix if suffix in {"csv", "json", "xlsx"} else "csv"


def parse_args() -> argparse.Namespace:
    root = repo_root()
    parser = argparse.ArgumentParser(
        description="Export Slack app UI strings missing from obj_stringtranslator."
    )
    parser.add_argument(
        "--languages",
        default=os.environ.get("TRANSLATION_EXPORT_LANGUAGES"),
        help=(
            "Comma-separated Slack locale codes to check. Defaults to the app's "
            "known Slack locale coverage set."
        ),
    )
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=root / "app",
        help="Python source directory to scan for _() calls.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parent / "output" / "missing_strings.csv",
        help="Output file path.",
    )
    parser.add_argument(
        "--format",
        choices=("csv", "json", "xlsx"),
        help="Output format. Defaults to the output file extension, then csv.",
    )
    parser.add_argument(
        "--include-english",
        action="store_true",
        help="Include English-like locales that runtime translation normally skips.",
    )
    parser.add_argument(
        "--split-by-language",
        action="store_true",
        help="Write one output file per resolved DB language.",
    )
    parser.add_argument(
        "--audience",
        choices=("internal", "translator"),
        default="internal",
        help=(
            "Output schema audience. 'translator' writes a single unnamed text "
            "column for vendor-facing per-language workbooks."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    root = repo_root()
    source_dir = args.source_dir.resolve()
    slack_locales = parse_languages(args.languages)
    language_map = fetch_language_map()
    targets = [
        resolve_locale_target(slack_locale, language_map)
        for slack_locale in slack_locales
    ]
    db_langs = [
        target.db_lang
        for target in targets
        if args.include_english or not target.is_english
    ]
    entries = collect_string_entries(source_dir, root=root)
    existing_labels = fetch_existing_labels(db_langs)
    rows = build_missing_rows(
        entries=entries,
        slack_locales=slack_locales,
        language_map=language_map,
        existing_labels_by_lang=existing_labels,
        include_english=args.include_english,
    )
    output_format = infer_format(args.output, args.format)
    if args.audience == "translator" and output_format == "json":
        raise ValueError("Translator exports support only csv and xlsx formats")
    if (
        args.audience == "translator"
        and not args.split_by_language
        and len(db_langs) != 1
    ):
        raise ValueError(
            "Translator exports must be split by language unless one target language is selected"
        )
    if args.split_by_language:
        output_paths = write_rows_by_language(
            rows,
            args.output,
            output_format,
            db_langs,
            args.audience,
        )
    else:
        write_rows(rows, args.output, output_format, args.audience)
        output_paths = [args.output]

    print(f"Scanned {len(entries)} unique strings from {source_dir}")
    print(f"Checked Slack locales: {', '.join(slack_locales)}")
    print(f"Exported {len(rows)} missing strings")
    for output_path in output_paths:
        print(f"  -> {output_path}")


if __name__ == "__main__":
    main()
