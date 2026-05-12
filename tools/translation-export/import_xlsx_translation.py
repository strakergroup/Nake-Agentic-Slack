"""Generate SQL inserts from a filled translation export workbook."""

from __future__ import annotations

import argparse
import csv
import glob
import re
import uuid
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

import openpyxl

LANG_COLUMN = "target_language"
LABEL_COLUMN = "source_text"
TRANSLATION_COLUMN = "target_text"
REQUIRED_COLUMNS = {LANG_COLUMN, LABEL_COLUMN, TRANSLATION_COLUMN}
TRANSLATOR_METADATA_SHEET = "_translation_metadata"
COLUMN_ALIASES = {
    LANG_COLUMN: ("db_lang",),
    LABEL_COLUMN: ("db_label",),
    TRANSLATION_COLUMN: ("translation",),
}
VALID_TAG_PATTERN = re.compile(r"<x id=(\d+)\s*/?>")
X_TAG_CANDIDATE_PATTERN = re.compile(r"</?x\b[^>]*>|<x\b[^>]*$")
ENGLISH_PREFIXES = ("en", "gb", "us")
LANGUAGE_FILENAME_PATTERN = re.compile(
    r"^(?:translations|missing_strings)_([^_]+)(?:_|$)"
)


@dataclass(frozen=True)
class TranslationValidationError:
    workbook: str
    row: int
    db_lang: str
    db_label: str
    translation: str
    error: str


def escape_sql_value(value: object) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\n", "\\n")
        .replace("\r", "\\r")
    )


def format_insert_statement(
    label: object,
    lang: str,
    translation: object,
    created: str,
    modified: str,
) -> str:
    uuid_str = str(uuid.uuid4())
    label_sql = escape_sql_value(label)
    lang_sql = escape_sql_value(lang)
    translation_sql = escape_sql_value(translation)
    return (
        "INSERT INTO `obj_stringtranslator` "
        "(`obj_uuid`, `created`, `modified`, `label`, `lang`, `langstring`, `active`) "
        f'VALUES ("{uuid_str}", "{created}", "{modified}", "{label_sql}", "{lang_sql}", "{translation_sql}", 1);\n'
    )


def format_delete_statement(lang: str, labels: list[str]) -> str:
    lang_sql = escape_sql_value(lang)
    labels_sql = ", ".join(f'"{escape_sql_value(label)}"' for label in labels)
    return (
        "DELETE FROM `obj_stringtranslator` "
        f'WHERE `lang` = "{lang_sql}" AND `label` IN ({labels_sql});\n'
    )


def find_tag_ids(text: str) -> set[str]:
    return set(VALID_TAG_PATTERN.findall(text))


def format_x_tag(tag_id: str) -> str:
    return f"<x id={tag_id}/>"


def find_malformed_tags(text: str) -> list[str]:
    candidates = X_TAG_CANDIDATE_PATTERN.findall(text)
    return [
        candidate
        for candidate in candidates
        if not VALID_TAG_PATTERN.fullmatch(candidate)
    ]


def validate_translation_tags(label: str, translation: str) -> list[str]:
    errors: list[str] = []
    malformed_label_tags = find_malformed_tags(label)
    malformed_translation_tags = find_malformed_tags(translation)
    if malformed_label_tags:
        errors.append(
            f"db_label contains malformed tag(s): {', '.join(malformed_label_tags)}"
        )
    if malformed_translation_tags:
        errors.append(
            "translation contains malformed tag(s): "
            + ", ".join(malformed_translation_tags)
        )

    expected_tag_ids = find_tag_ids(label)
    translation_tag_ids = find_tag_ids(translation)
    missing_tag_ids = sorted(expected_tag_ids - translation_tag_ids, key=int)
    extra_tag_ids = sorted(translation_tag_ids - expected_tag_ids, key=int)
    if missing_tag_ids:
        errors.append(
            "translation is missing placeholder tag(s): "
            + ", ".join(format_x_tag(tag_id) for tag_id in missing_tag_ids)
        )
    if extra_tag_ids:
        errors.append(
            "translation contains unexpected placeholder tag(s): "
            + ", ".join(format_x_tag(tag_id) for tag_id in extra_tag_ids)
        )

    return errors


def validate_translation_text(lang: str, label: str, translation: str) -> list[str]:
    if lang.strip().lower().startswith(ENGLISH_PREFIXES):
        return []
    if label.strip() == translation.strip():
        return [
            "translation matches source text for non-English language; review MT output"
        ]
    return []


def has_suspicious_matching_source_batch(total_rows: int, matching_rows: int) -> bool:
    if total_rows < 5:
        return False
    return matching_rows / total_rows >= 0.8


def header_indexes(sheet) -> dict[str, int]:
    headers = [cell.value for cell in sheet[1]]
    return {
        str(header): index + 1
        for index, header in enumerate(headers)
        if isinstance(header, str)
    }


def column_index(indexes: dict[str, int], column: str) -> int | None:
    if column in indexes:
        return indexes[column]
    return next(
        (
            indexes[alias]
            for alias in COLUMN_ALIASES.get(column, ())
            if alias in indexes
        ),
        None,
    )


def infer_target_language_from_path(workbook_path: Path) -> str | None:
    """Infer DB language from names like translations_fr_updated__French.xlsx."""
    stem = workbook_path.stem
    match = LANGUAGE_FILENAME_PATTERN.match(stem)
    if match:
        return match.group(1) or None
    if "_" not in stem:
        return None
    return stem.rsplit("_", 1)[1] or None


def read_single_column_translation(sheet, row_number: int) -> object:
    for column in range(2, sheet.max_column + 1):
        value = sheet.cell(row=row_number, column=column).value
        if value is not None and str(value).strip():
            return value
    return sheet.cell(row=row_number, column=1).value


def read_translator_metadata(workbook) -> dict[int, str]:
    if TRANSLATOR_METADATA_SHEET not in workbook.sheetnames:
        return {}
    metadata_sheet = workbook[TRANSLATOR_METADATA_SHEET]
    indexes = header_indexes(metadata_sheet)
    label_index = column_index(indexes, LABEL_COLUMN)
    if label_index is None:
        return {}
    labels_by_visible_row: dict[int, str] = {}
    for metadata_row_number in range(2, metadata_sheet.max_row + 1):
        label = metadata_sheet.cell(
            row=metadata_row_number,
            column=label_index,
        ).value
        if label:
            labels_by_visible_row[metadata_row_number - 1] = str(label)
    return labels_by_visible_row


def collect_insert_statements(
    workbook_path: Path,
    created: str,
    modified: str,
) -> tuple[list[str], dict[str, set[str]], list[TranslationValidationError]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        indexes = header_indexes(sheet)
        label_index = column_index(indexes, LABEL_COLUMN)
        lang_index = column_index(indexes, LANG_COLUMN)
        translation_index = column_index(indexes, TRANSLATION_COLUMN)
        translator_metadata = read_translator_metadata(workbook)
        is_single_column_translator_workbook = bool(translator_metadata)
        inferred_lang = (
            None
            if lang_index is not None
            else infer_target_language_from_path(workbook_path)
        )
        missing_columns = []
        if not is_single_column_translator_workbook:
            missing_columns = [
                column
                for column, index in (
                    (LABEL_COLUMN, label_index),
                    (TRANSLATION_COLUMN, translation_index),
                )
                if index is None
            ]
        if lang_index is None and inferred_lang is None:
            missing_columns.append(LANG_COLUMN)
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise ValueError(f"Workbook is missing required columns: {missing}")
        if not is_single_column_translator_workbook:
            assert label_index is not None
            assert translation_index is not None

        statements: list[str] = []
        seen_pairs: set[tuple[str, str]] = set()
        delete_labels_by_lang: dict[str, set[str]] = defaultdict(set)
        validation_errors: list[TranslationValidationError] = []
        row_counts_by_lang: dict[str, int] = defaultdict(int)
        matching_source_rows_by_lang: dict[str, list[TranslationValidationError]] = (
            defaultdict(list)
        )
        start_row = 1 if is_single_column_translator_workbook else 2
        for row_number in range(start_row, sheet.max_row + 1):
            label = (
                translator_metadata.get(row_number)
                if is_single_column_translator_workbook
                else sheet.cell(row=row_number, column=label_index).value
            )
            lang = (
                sheet.cell(row=row_number, column=lang_index).value
                if lang_index is not None
                else inferred_lang
            )
            translation = (
                read_single_column_translation(sheet, row_number)
                if is_single_column_translator_workbook
                else sheet.cell(row=row_number, column=translation_index).value
            )
            if not label or not lang or not translation or not str(translation).strip():
                continue
            label_text = str(label)
            lang_text = str(lang)
            translation_text = str(translation)
            if (
                is_single_column_translator_workbook
                and translation_text.strip() == label_text.strip()
            ):
                continue
            row_errors = validate_translation_tags(label_text, translation_text)
            if not lang_text.strip().lower().startswith(ENGLISH_PREFIXES):
                row_counts_by_lang[lang_text] += 1
                if label_text.strip() == translation_text.strip():
                    matching_source_rows_by_lang[lang_text].append(
                        TranslationValidationError(
                            workbook=str(workbook_path),
                            row=row_number,
                            db_lang=lang_text,
                            db_label=label_text,
                            translation=translation_text,
                            error=(
                                "translation matches source text for non-English "
                                "language across a suspicious share of this workbook; "
                                "review MT output"
                            ),
                        )
                    )
            validation_errors.extend(
                TranslationValidationError(
                    workbook=str(workbook_path),
                    row=row_number,
                    db_lang=lang_text,
                    db_label=label_text,
                    translation=translation_text,
                    error=error,
                )
                for error in row_errors
            )
            if row_errors:
                continue
            pair = (lang_text, label_text)
            if pair in seen_pairs:
                continue
            seen_pairs.add(pair)
            delete_labels_by_lang[lang_text].add(label_text)
            statements.append(
                format_insert_statement(
                    label_text,
                    lang_text,
                    translation_text,
                    created,
                    modified,
                )
            )
        for lang, matching_rows in matching_source_rows_by_lang.items():
            if has_suspicious_matching_source_batch(
                row_counts_by_lang[lang], len(matching_rows)
            ):
                validation_errors.extend(matching_rows)
        return statements, delete_labels_by_lang, validation_errors
    finally:
        workbook.close()


def collect_refresh_delete_statements(
    delete_labels_by_lang: dict[str, set[str]],
) -> list[str]:
    statements: list[str] = []
    for lang in sorted(delete_labels_by_lang):
        labels = sorted(delete_labels_by_lang[lang])
        if labels:
            statements.append(format_delete_statement(lang, labels))
    return statements


def resolve_workbook_paths(input_pattern: str) -> list[Path]:
    matches = sorted(glob.glob(input_pattern))
    if matches:
        return [Path(match) for match in matches]
    return [Path(input_pattern)]


def default_validation_report_path(output_path: Path) -> Path:
    return output_path.with_name(f"{output_path.stem}_validation_errors.csv")


def write_validation_report(
    validation_errors: list[TranslationValidationError],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as report_file:
        writer = csv.DictWriter(
            report_file,
            fieldnames=[
                "workbook",
                "row",
                "db_lang",
                "db_label",
                "translation",
                "error",
            ],
        )
        writer.writeheader()
        writer.writerows(asdict(error) for error in validation_errors)


def write_import_sql(
    workbook_paths: list[Path],
    output_path: Path,
    validation_report_path: Path | None = None,
    refresh_delete: bool = True,
) -> int:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    statements: list[str] = []
    delete_labels_by_lang: dict[str, set[str]] = defaultdict(set)
    validation_errors: list[TranslationValidationError] = []
    validation_report_path = validation_report_path or default_validation_report_path(
        output_path
    )
    for workbook_path in workbook_paths:
        workbook_statements, workbook_delete_labels_by_lang, workbook_errors = (
            collect_insert_statements(
                workbook_path,
                timestamp,
                timestamp,
            )
        )
        for lang, labels in workbook_delete_labels_by_lang.items():
            delete_labels_by_lang[lang].update(labels)
        statements.extend(workbook_statements)
        validation_errors.extend(workbook_errors)

    if validation_errors:
        write_validation_report(validation_errors, validation_report_path)
        raise ValueError(
            f"Validation failed for {len(validation_errors)} translation row(s). "
            f"See {validation_report_path}."
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as sql_file:
        if refresh_delete:
            sql_file.writelines(
                collect_refresh_delete_statements(delete_labels_by_lang)
            )
        sql_file.writelines(statements)
    return len(statements)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate obj_stringtranslator SQL from a filled workbook."
    )
    parser.add_argument(
        "--input",
        default=str(Path(__file__).parent / "output" / "missing_strings.xlsx"),
        help="Filled workbook generated by export_missing_strings.py. Globs are supported.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parent / "output" / "import.sql",
        help="SQL output path.",
    )
    parser.add_argument(
        "--validation-report",
        type=Path,
        help="CSV output path for placeholder validation errors.",
    )
    parser.add_argument(
        "--refresh-delete",
        action=argparse.BooleanOptionalAction,
        default=True,
        help=(
            "Delete existing obj_stringtranslator rows for the exact lang/label "
            "pairs in the input before inserting. Enabled by default."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    count = write_import_sql(
        resolve_workbook_paths(args.input),
        args.output,
        args.validation_report,
        args.refresh_delete,
    )
    print(f"Generated {count} insert statements -> {args.output}")


if __name__ == "__main__":
    main()
