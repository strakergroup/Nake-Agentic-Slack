import importlib.util
import sys
from pathlib import Path

import openpyxl
import pytest

TOOLS_DIR = Path(__file__).resolve().parents[2] / "tools" / "translation-export"
APP_ROOT = Path(__file__).resolve().parents[2]
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))


def load_tool_module(module_name: str):
    module_path = TOOLS_DIR / f"{module_name}.py"
    spec = importlib.util.spec_from_file_location(module_name, module_path)
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


MT_FILL = load_tool_module("mt_fill_translations")
IMPORT_SQL = load_tool_module("import_xlsx_translation")
EXPORT = load_tool_module("export_missing_strings")


class FakeGoogleTranslation:
    def __init__(self, translated_text: str):
        self.translated_text = translated_text


class FakeGoogleResponse:
    def __init__(self, translations: list[str]):
        self.translations = [
            FakeGoogleTranslation(translation) for translation in translations
        ]


class RecordingGoogleClient:
    def __init__(self, translate_callback=None):
        self.calls = []
        self.translate_callback = translate_callback or (
            lambda contents, target_language_code: [
                f"{target_language_code}:{content}" for content in contents
            ]
        )

    def translate_text(self, **kwargs):
        self.calls.append(kwargs)
        return FakeGoogleResponse(
            self.translate_callback(
                kwargs["contents"],
                kwargs["target_language_code"],
            )
        )


class FakeTranslationResult:
    def __init__(self, row):
        self.row = row

    def fetchone(self):
        return self.row


class FakeConnection:
    def __init__(self, translations_by_label):
        self.translations_by_label = translations_by_label
        self.queries = []

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return False

    def execute(self, statement):
        if "input" not in statement._bindparams:
            return []
        label = statement._bindparams["input"].value
        self.queries.append(label)
        translation = self.translations_by_label.get(label)
        return FakeTranslationResult(None if translation is None else (translation,))


class FakeEngine:
    def __init__(self, translations_by_label):
        self.connection = FakeConnection(translations_by_label)

    def connect(self):
        return self.connection


@pytest.fixture(autouse=True)
def google_mt_env(monkeypatch):
    monkeypatch.setenv("GOOGLE_CLOUD_PROJECT", "translation-export-test")
    monkeypatch.delenv("GOOGLE_CLOUD_LOCATION", raising=False)
    monkeypatch.delenv("TRANSLATION_EXPORT_MT_BATCH_SIZE", raising=False)


def make_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(["en", "fr", "Submit <x id=1/>", "", 0])
    sheet.append(["en", "de", "Cancel", "Abbrechen", 0])
    workbook.save(path)
    workbook.close()


def make_legacy_translator_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["source_text", "translation", "notes"])
    sheet.append(["Submit <x id=1>", "", "Preserve <x id=N> placeholder tags."])
    sheet.append(["Cancel", "Annuler", ""])
    workbook.save(path)
    workbook.close()


def test_write_translator_xlsx_uses_single_unnamed_column_with_metadata(tmp_path):
    workbook_path = tmp_path / "missing_strings_fr.xlsx"
    rows = [
        EXPORT.MissingStringRow(
            source_language="en",
            target_language="fr",
            source_text="Submit <x id=1/>",
            target_text="",
            max_length=20,
        )
    ]

    EXPORT.write_rows(rows, workbook_path, "xlsx", audience="translator")

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert sheet.title == "Translations"
        assert sheet.max_column == 1
        assert sheet["A1"].value == "Submit <x id=1/>"
        metadata_sheet = workbook[EXPORT.TRANSLATOR_METADATA_SHEET]
        assert metadata_sheet.sheet_state == "hidden"
        assert [cell.value for cell in metadata_sheet[1]] == [
            "source_text",
            "max_length",
        ]
        assert metadata_sheet["A2"].value == "Submit <x id=1/>"
        assert metadata_sheet["B2"].value == 20
    finally:
        workbook.close()


def test_fill_workbook_translations_populates_blank_translation(tmp_path, monkeypatch):
    workbook_path = tmp_path / "missing_strings.xlsx"
    make_workbook(workbook_path)
    client = RecordingGoogleClient()
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {})

    filled, total = MT_FILL.fill_workbook_translations(
        workbook_path,
        google_client=client,
    )

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert filled == 1
        assert total == 2
        assert sheet["D2"].value == "fr:Submit <x id=1/>"
        assert sheet["D3"].value == "Abbrechen"
        assert client.calls[0]["contents"] == ["Submit <x id=1/>"]
        assert client.calls[0]["target_language_code"] == "fr"
        assert client.calls[0]["mime_type"] == "text/html"
    finally:
        workbook.close()


def test_fill_legacy_translator_workbook_infers_language_from_filename(
    tmp_path, monkeypatch
):
    workbook_path = tmp_path / "translations_fr.xlsx"
    make_legacy_translator_workbook(workbook_path)
    client = RecordingGoogleClient()
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {})

    filled, total = MT_FILL.fill_workbook_translations(
        workbook_path,
        google_client=client,
    )

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert filled == 1
        assert total == 2
        assert sheet["B2"].value == "fr:Submit <x id=1>"
        assert sheet["B3"].value == "Annuler"
        assert client.calls[0]["contents"] == ["Submit <x id=1>"]
        assert client.calls[0]["target_language_code"] == "fr"
    finally:
        workbook.close()


def test_fill_single_column_translator_workbook_replaces_visible_source_text(
    tmp_path, monkeypatch
):
    workbook_path = tmp_path / "translations_fr.xlsx"
    rows = [
        EXPORT.MissingStringRow(
            source_language="en",
            target_language="fr",
            source_text="Submit <x id=1/>",
            target_text="",
            max_length=0,
        )
    ]
    EXPORT.write_translator_xlsx(rows, workbook_path)
    client = RecordingGoogleClient()
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {})

    filled, total = MT_FILL.fill_workbook_translations(
        workbook_path,
        google_client=client,
    )

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert filled == 1
        assert total == 1
        assert sheet.max_column == 1
        assert sheet["A1"].value == "fr:Submit <x id=1/>"
        assert client.calls[0]["contents"] == ["Submit <x id=1/>"]
        assert client.calls[0]["target_language_code"] == "fr"
    finally:
        workbook.close()


def test_fill_workbook_translations_decodes_html_entities_and_preserves_tags(
    tmp_path, monkeypatch
):
    workbook_path = tmp_path / "missing_strings.xlsx"
    make_workbook(workbook_path)
    client = RecordingGoogleClient(
        lambda contents, target_language_code: [
            "L&#39;envoi &quot;OK&quot; &amp; <x id=1/>&nbsp;" for _ in contents
        ]
    )
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {})

    filled, total = MT_FILL.fill_workbook_translations(
        workbook_path,
        google_client=client,
    )

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert filled == 1
        assert total == 2
        assert sheet["D2"].value == 'L\'envoi "OK" & <x id=1/>\xa0'
    finally:
        workbook.close()


def test_fill_workbook_maps_korean_db_lang_to_mt_code(tmp_path, monkeypatch):
    workbook_path = tmp_path / "missing_strings_kr.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(["en", "kr", "Submit", "", 0])
    workbook.save(workbook_path)
    workbook.close()

    client = RecordingGoogleClient(lambda contents, target_language_code: ["제출"])
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {"kr": "ko"})

    filled, total = MT_FILL.fill_workbook_translations(
        workbook_path,
        google_client=client,
    )

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert filled == 1
        assert total == 1
        assert [call["target_language_code"] for call in client.calls] == ["ko"]
        assert sheet["B2"].value == "kr"
        assert sheet["D2"].value == "제출"
    finally:
        workbook.close()


def test_validate_mt_translation_rejects_unchanged_non_english_output():
    with pytest.raises(ValueError, match="unchanged source text"):
        MT_FILL.validate_mt_translation("Submit", "kr", "Submit")


def test_fill_workbook_raises_and_notes_unchanged_non_english_output(
    tmp_path, monkeypatch
):
    workbook_path = tmp_path / "missing_strings.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    for label in ["Submit", "Cancel", "Done", "Back", "Next"]:
        sheet.append(["en", "kr", label, "", 0])
    workbook.save(workbook_path)
    workbook.close()
    client = RecordingGoogleClient(lambda contents, target_language_code: contents)
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {})

    with pytest.raises(ValueError, match="MT fill failed for 1 row"):
        MT_FILL.fill_workbook_translations(workbook_path, google_client=client)

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert sheet["F1"].value == "notes"
        assert "suspicious share" in sheet["F2"].value
        assert sheet["D2"].value == "Submit"
    finally:
        workbook.close()


def test_write_import_sql_skips_empty_translations(tmp_path):
    workbook_path = tmp_path / "missing_strings.xlsx"
    output_path = tmp_path / "import.sql"
    make_workbook(workbook_path)

    count = IMPORT_SQL.write_import_sql([workbook_path], output_path)

    sql = output_path.read_text(encoding="utf-8")
    assert count == 1
    assert '"de"' in sql
    assert '"Cancel"' in sql
    assert '"Abbrechen"' in sql
    assert "Submit" not in sql


def test_write_import_sql_accepts_legacy_translator_workbook(tmp_path):
    workbook_path = tmp_path / "translations_fr.xlsx"
    output_path = tmp_path / "import.sql"
    make_legacy_translator_workbook(workbook_path)

    count = IMPORT_SQL.write_import_sql([workbook_path], output_path)

    sql = output_path.read_text(encoding="utf-8")
    assert count == 1
    assert '"fr"' in sql
    assert '"Cancel"' in sql
    assert '"Annuler"' in sql
    assert "Submit" not in sql


def test_workbook_language_inference_reads_db_code_before_update_suffix():
    assert (
        IMPORT_SQL.infer_target_language_from_path(
            Path("translations_fr-ca_updated__French_Canada.xlsx")
        )
        == "fr-ca"
    )
    assert (
        IMPORT_SQL.infer_target_language_from_path(
            Path("translations_jp_updated__Japanese.xlsx")
        )
        == "jp"
    )
    assert MT_FILL.infer_target_language_from_path(Path("translations_fr.xlsx")) == "fr"
    assert (
        MT_FILL.infer_target_language_from_path(Path("missing_strings_zh-CN.xlsx"))
        == "zh-CN"
    )


def test_write_import_sql_accepts_single_column_translator_workbook(tmp_path):
    workbook_path = tmp_path / "translations_fr.xlsx"
    output_path = tmp_path / "import.sql"
    rows = [
        EXPORT.MissingStringRow(
            source_language="en",
            target_language="fr",
            source_text="Submit <x id=1/>",
            target_text="",
            max_length=0,
        )
    ]
    EXPORT.write_translator_xlsx(rows, workbook_path)

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        sheet["A1"].value = "Soumettre <x id=1/>"
        workbook.save(workbook_path)
    finally:
        workbook.close()

    count = IMPORT_SQL.write_import_sql([workbook_path], output_path)

    sql = output_path.read_text(encoding="utf-8")
    assert count == 1
    assert '"fr"' in sql
    assert '"Submit <x id=1/>"' in sql
    assert '"Soumettre <x id=1/>"' in sql


def test_write_import_sql_reads_returned_translator_second_column(tmp_path):
    workbook_path = tmp_path / "translations_fr-ca_updated__French_Canada.xlsx"
    output_path = tmp_path / "import.sql"
    rows = [
        EXPORT.MissingStringRow(
            source_language="en",
            target_language="fr-ca",
            source_text="Submit <x id=1/>",
            target_text="",
            max_length=0,
        )
    ]
    EXPORT.write_translator_xlsx(rows, workbook_path)

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        sheet["B1"].value = "Soumettre <x id=1/>"
        workbook.save(workbook_path)
    finally:
        workbook.close()

    count = IMPORT_SQL.write_import_sql([workbook_path], output_path)

    sql = output_path.read_text(encoding="utf-8")
    assert count == 1
    assert '"fr-ca"' in sql
    assert '"French_Canada"' not in sql
    assert '"Submit <x id=1/>"' in sql
    assert '"Soumettre <x id=1/>"' in sql


def test_write_import_sql_deletes_only_generated_lang_label_pairs(tmp_path):
    workbook_path = tmp_path / "missing_strings.xlsx"
    output_path = tmp_path / "import.sql"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(["en", "fr", 'Quote "label"', "Libellé", 0])
    sheet.append(["en", "fr-ca", "Regional label", "Libellé régional", 0])
    sheet.append(["en", "de", "Needs MT", "", 0])
    workbook.save(workbook_path)
    workbook.close()

    count = IMPORT_SQL.write_import_sql([workbook_path], output_path)

    sql = output_path.read_text(encoding="utf-8")
    assert count == 2
    assert sql.count("DELETE FROM `obj_stringtranslator`") == 2
    assert (
        'DELETE FROM `obj_stringtranslator` WHERE `lang` = "fr" '
        'AND `label` IN ("Quote \\"label\\"");'
    ) in sql
    assert (
        'DELETE FROM `obj_stringtranslator` WHERE `lang` = "fr-ca" '
        'AND `label` IN ("Regional label");'
    ) in sql
    assert 'WHERE `lang` = "de"' not in sql
    assert "DELETE FROM `obj_stringtranslator` WHERE `lang` IN" not in sql
    assert sql.count("INSERT INTO `obj_stringtranslator`") == 2
    assert sql.index("DELETE FROM `obj_stringtranslator`") < sql.index(
        "INSERT INTO `obj_stringtranslator`"
    )


def test_write_import_sql_can_disable_refresh_deletes(tmp_path):
    workbook_path = tmp_path / "missing_strings.xlsx"
    output_path = tmp_path / "import.sql"
    make_workbook(workbook_path)

    count = IMPORT_SQL.write_import_sql(
        [workbook_path],
        output_path,
        refresh_delete=False,
    )

    sql = output_path.read_text(encoding="utf-8")
    assert count == 1
    assert "DELETE FROM `obj_stringtranslator`" not in sql
    assert sql.count("INSERT INTO `obj_stringtranslator`") == 1


def test_write_import_sql_deduplicates_generated_lang_label_pairs(tmp_path):
    workbook_path = tmp_path / "missing_strings.xlsx"
    output_path = tmp_path / "import.sql"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(["en", "fr", "Submit", "Soumettre", 0])
    sheet.append(["en", "fr", "Submit", "Soumettre", 20])
    workbook.save(workbook_path)
    workbook.close()

    count = IMPORT_SQL.write_import_sql([workbook_path], output_path)

    sql = output_path.read_text(encoding="utf-8")
    assert count == 1
    assert sql.count("DELETE FROM `obj_stringtranslator`") == 1
    assert sql.count("INSERT INTO `obj_stringtranslator`") == 1


def test_validate_translation_tags_reports_missing_and_bad_tags():
    errors = IMPORT_SQL.validate_translation_tags(
        "Submit <x id=1/>",
        "Soumettre <x id=two>",
    )

    assert "translation contains malformed tag(s): <x id=two>" in errors
    assert "translation is missing placeholder tag(s): <x id=1/>" in errors


def test_validate_translation_tags_accepts_self_closing_and_legacy_tags():
    assert (
        IMPORT_SQL.validate_translation_tags(
            "Submit <x id=1/>",
            "Soumettre <x id=1/>",
        )
        == []
    )
    assert (
        IMPORT_SQL.validate_translation_tags(
            "Submit <x id=1>",
            "Soumettre <x id=1>",
        )
        == []
    )


def test_runtime_translator_restores_self_closing_and_legacy_x_tags(monkeypatch):
    import types

    fake_engine = FakeEngine(
        {
            "Hello <x id=1/>": "Bonjour <x id=1/>",
            "Legacy <x id=1>": "Héritage <x id=1>",
        }
    )
    fake_database = types.SimpleNamespace(
        engines={
            "translators_readonly": FakeEngine({}),
            "sitemanager_readonly": fake_engine,
        }
    )
    fake_buglog = types.SimpleNamespace(notify_exception=lambda exc: None)
    monkeypatch.setitem(sys.modules, "app.database", fake_database)
    monkeypatch.setitem(sys.modules, "app.slack.buglog_notifier", fake_buglog)
    module_name = "app.translate_runtime_test"
    sys.modules.pop(module_name, None)
    spec = importlib.util.spec_from_file_location(
        module_name,
        APP_ROOT / "app" / "translate.py",
    )
    assert spec is not None
    assert spec.loader is not None
    translate_module = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, module_name, translate_module)
    spec.loader.exec_module(translate_module)
    translate_module.Translator.BCP_47_TO_SHORTNAME = {"fr-FR": "fr"}

    translator = translate_module.Translator("fr-FR")

    assert translator.translate("Hello {name}") == ("Bonjour {name}", True)
    assert translator.translate("Legacy {name}") == ("Héritage {name}", True)
    assert fake_engine.connection.queries == [
        "Hello <x id=1/>",
        "Legacy <x id=1/>",
        "Legacy <x id=1>",
    ]


def test_validate_translation_text_reports_matching_non_english_source():
    errors = IMPORT_SQL.validate_translation_text("kr", "Submit", "Submit")

    assert (
        "translation matches source text for non-English language; review MT output"
        in errors
    )
    assert IMPORT_SQL.validate_translation_text("en", "Submit", "Submit") == []


def test_write_import_sql_fails_and_reports_placeholder_validation_errors(tmp_path):
    workbook_path = tmp_path / "missing_strings.xlsx"
    output_path = tmp_path / "import.sql"
    report_path = tmp_path / "validation_errors.csv"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(
        [
            "en",
            "fr",
            "Submit <x id=1/>",
            "Soumettre",
            0,
        ]
    )
    workbook.save(workbook_path)
    workbook.close()

    try:
        IMPORT_SQL.write_import_sql([workbook_path], output_path, report_path)
    except ValueError as exc:
        assert "Validation failed for 1 translation row(s)" in str(exc)
    else:
        raise AssertionError("Expected placeholder validation to fail")

    report = report_path.read_text(encoding="utf-8")
    assert "translation is missing placeholder tag(s): <x id=1/>" in report
    assert not output_path.exists()


def test_write_import_sql_fails_on_matching_non_english_source_text(tmp_path):
    workbook_path = tmp_path / "missing_strings.xlsx"
    output_path = tmp_path / "import.sql"
    report_path = tmp_path / "validation_errors.csv"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    for label in ["Submit", "Cancel", "Done", "Back", "Next"]:
        sheet.append(["en", "kr", label, label, 0])
    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(ValueError, match="Validation failed for 5 translation row"):
        IMPORT_SQL.write_import_sql([workbook_path], output_path, report_path)

    report = report_path.read_text(encoding="utf-8")
    assert "suspicious share of this workbook" in report
    assert not output_path.exists()


def test_fill_workbooks_accepts_glob_input(tmp_path):
    first_path = tmp_path / "missing_strings_fr.xlsx"
    second_path = tmp_path / "missing_strings_de.xlsx"
    make_workbook(first_path)
    make_workbook(second_path)

    paths = MT_FILL.resolve_workbook_paths(str(tmp_path / "missing_strings_*.xlsx"))

    assert paths == [second_path, first_path]


def test_format_insert_statement_escapes_sql_values():
    statement = IMPORT_SQL.format_insert_statement(
        'Label "quoted"',
        "fr",
        "Line 1\nLine 2",
        "2026-04-30 00:00:00",
        "2026-04-30 00:00:00",
    )

    assert 'Label \\"quoted\\"' in statement
    assert "Line 1\\nLine 2" in statement


def test_google_translate_texts_uses_google_v3_batch_request():
    client = RecordingGoogleClient(lambda contents, target_language_code: ["Bonjour"])

    translations = MT_FILL.google_translate_texts(
        ["Hello <x id=1/>"],
        "fr",
        client=client,
        project_id="project-123",
        location="global",
    )

    assert translations == ["Bonjour"]
    assert client.calls == [
        {
            "contents": ["Hello <x id=1/>"],
            "parent": "projects/project-123/locations/global",
            "mime_type": "text/html",
            "source_language_code": "en",
            "target_language_code": "fr",
        }
    ]


def test_google_translate_texts_requires_matching_response_count():
    client = RecordingGoogleClient(lambda contents, target_language_code: ["Bonjour"])

    with pytest.raises(ValueError, match="1 translation"):
        MT_FILL.google_translate_texts(
            ["Hello", "Bye"],
            "fr",
            client=client,
            project_id="project-123",
            location="global",
        )


def test_fill_workbook_batches_rows_by_google_language(tmp_path, monkeypatch):
    workbook_path = tmp_path / "missing_strings.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(["en", "fr", "Submit", "", 0])
    sheet.append(["en", "fr", "Cancel", "", 0])
    sheet.append(["en", "fr", "Done", "", 0])
    sheet.append(["en", "de", "Back", "", 0])
    workbook.save(workbook_path)
    workbook.close()

    client = RecordingGoogleClient(
        lambda contents, target_language_code: [
            f"{target_language_code}:{content}" for content in contents
        ]
    )
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {})

    filled, total = MT_FILL.fill_workbook_translations(
        workbook_path,
        batch_size=2,
        google_client=client,
    )

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook.active
        assert filled == 4
        assert total == 4
        assert sheet["D2"].value == "fr:Submit"
        assert sheet["D3"].value == "fr:Cancel"
        assert sheet["D4"].value == "fr:Done"
        assert sheet["D5"].value == "de:Back"
    finally:
        workbook.close()

    assert [
        (call["target_language_code"], call["contents"]) for call in client.calls
    ] == [
        ("fr", ["Submit", "Cancel"]),
        ("fr", ["Done"]),
        ("de", ["Back"]),
    ]


def test_fill_workbook_rejects_non_english_db_lang_mapped_to_english(
    tmp_path, monkeypatch
):
    workbook_path = tmp_path / "missing_strings.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "source_language",
            "target_language",
            "source_text",
            "target_text",
            "max_length",
        ]
    )
    sheet.append(["en", "kr", "Submit", "", 0])
    workbook.save(workbook_path)
    workbook.close()
    client = RecordingGoogleClient()
    monkeypatch.setattr(MT_FILL, "fetch_mt_language_map", lambda: {"kr": "en"})

    with pytest.raises(ValueError, match="English fallback"):
        MT_FILL.fill_workbook_translations(workbook_path, google_client=client)

    assert client.calls == []


def test_configured_google_project_requires_env(monkeypatch):
    monkeypatch.delenv("GOOGLE_CLOUD_PROJECT", raising=False)

    with pytest.raises(ValueError, match="GOOGLE_CLOUD_PROJECT"):
        MT_FILL.configured_google_project()


def test_configured_batch_size_reads_env(monkeypatch):
    monkeypatch.setenv("TRANSLATION_EXPORT_MT_BATCH_SIZE", "25")

    assert MT_FILL.configured_batch_size() == 25
